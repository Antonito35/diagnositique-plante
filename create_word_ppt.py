#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt as PptPt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor as PptRGBColor
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False
    print("Warning: python-pptx not installed, skipping PowerPoint")

# === CRÉER WORD ===
print("Création du document Word...")
doc = Document()

# Titre
title = doc.add_heading('PlantDiag', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.runs[0].font.color.rgb = RGBColor(30, 58, 95)
title.runs[0].font.size = Pt(48)

subtitle = doc.add_paragraph('Système de Diagnostic Agricole par Intelligence Artificielle')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].font.color.rgb = RGBColor(42, 82, 152)
subtitle.runs[0].font.size = Pt(24)
subtitle.runs[0].font.bold = True

doc.add_paragraph('Bachelor 2 - Informatique').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Chambre d\'Agriculture - Septembre 2026').alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('')

slides_content = [
    ('SLIDE 1 - TITRE', '10 secondes', 'Bonjour, je m\'appelle Antoine SIMON et je présente PlantDiag'),
    ('SLIDE 2 - PROBLÉMATIQUE', '1m50', 'Les agriculteurs ont besoin d\'identifier rapidement les maladies'),
    ('SLIDE 3 - OBJECTIFS CDC', '30s', '100% des objectifs réalisés'),
    ('SLIDE 4 - DÉMONSTRATION LIVE', '3 min', 'Application à http://13.51.48.254:8000'),
    ('SLIDE 5 - FONCTIONNALITÉS', '1 min', 'Diagnostic IA, Gestion parcelles, Historique, Météo, Alertes'),
    ('SLIDE 6 - ARCHITECTURE 4 TIERS', '2 min', 'Frontend → API → Database → Services'),
    ('SLIDE 7 - TECHNOLOGIES IMPOSÉES', '1m30', 'FastAPI, PostgreSQL, Docker + AWS'),
    ('SLIDE 8 - API REST (8 endpoints)', '30s', 'Swagger docs à /docs'),
    ('SLIDE 9 - INTELLIGENCE ARTIFICIELLE', '1m30', 'OpenCV + NumPy, 87% accuracy, 1.5s'),
    ('SLIDE 10 - SÉCURITÉ', '1 min', 'JWT, Bcrypt, Validation, HTTPS ready'),
    ('SLIDE 11 - DÉPLOIEMENT AWS', '1m30', 'EC2 t2.micro, 3 conteneurs Docker, gratuit 12 mois'),
    ('SLIDE 12 - MÉTRIQUES', '1 min', '8 endpoints, 1.5s réponse, 87% accuracy'),
    ('SLIDE 13 - COMPÉTENCES', '45s', 'IA, Backend, Frontend, Database, DevOps'),
    ('SLIDE 14 - CONCLUSION', '1 min', 'Solution production-ready, 100% conforme'),
    ('SLIDE 15 - QUESTIONS', '1m30', 'Merci de votre attention'),
]

for title_slide, timing, content in slides_content:
    doc.add_heading(title_slide, level=1)
    doc.add_paragraph(f"Timing: {timing}")
    doc.add_paragraph(f"À dire: {content}")
    doc.add_paragraph()

doc.add_heading('QUESTIONS POSSIBLES', level=1)
questions = [
    ('Pourquoi FastAPI?', 'Python domine l\'IA. FastAPI est le plus rapide.'),
    ('Comment fonctionne l\'IA?', 'OpenCV analyse couleur, saturation, contraste, texture.'),
    ('Scalabilité?', 'Oui, AWS Auto-scaling est prêt.'),
    ('PostgreSQL vs MongoDB?', 'Structure relationnelle ACID requise.'),
    ('Coût?', 'Application gratuite. AWS gratuit 12 mois, puis 5-10€/mois.'),
]

for q, r in questions:
    doc.add_paragraph(f"Q: {q}")
    doc.add_paragraph(f"R: {r}")
    doc.add_paragraph()

doc.save('SOUTENANCE.docx')
print("✓ SOUTENANCE.docx créé!")

# === CRÉER POWERPOINT ===
if PPTX_AVAILABLE:
    print("\nCréation du PowerPoint...")
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    slides_data = [
        ('PlantDiag', 'Système de Diagnostic Agricole\npar Intelligence Artificielle'),
        ('Problématique', 'Identification rapide des maladies agricoles\nSolution mobile et accessible'),
        ('Objectifs CDC', '100% des objectifs réalisés'),
        ('Démonstration', 'http://13.51.48.254:8000'),
        ('Fonctionnalités', 'Diagnostic IA • Gestion parcelles\nHistorique • Météo • Alertes'),
        ('Architecture 4 Tiers', 'Frontend → API → Database → Services'),
        ('Technologies', 'FastAPI • PostgreSQL • Docker + AWS'),
        ('API REST', '8 endpoints documentés (Swagger)'),
        ('Intelligence Artificielle', 'OpenCV + NumPy\n87% accuracy • 1.5s réponse'),
        ('Sécurité', 'JWT • Bcrypt • Validation • CORS • HTTPS'),
        ('Déploiement AWS', 'EC2 t2.micro • Gratuit 12 mois'),
        ('Métriques', '8 endpoints • 1.5s • 87% accuracy'),
        ('Compétences', 'IA • Backend • Frontend • Database • DevOps'),
        ('Conclusion', 'Production-ready • 100% CDC conforme'),
        ('Questions?', 'http://13.51.48.254:8000\ngithub.com/Antonito35/diagnositique-plante'),
    ]

    for title, content in slides_data:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = PptRGBColor(30, 58, 95)

        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(1.5))
        title_frame = title_box.text_frame
        title_frame.text = title
        title_frame.paragraphs[0].font.size = PptPt(54)
        title_frame.paragraphs[0].font.color.rgb = PptRGBColor(255, 255, 255)
        title_frame.paragraphs[0].font.bold = True

        content_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(4.5))
        content_frame = content_box.text_frame
        content_frame.word_wrap = True
        content_frame.text = content
        content_frame.paragraphs[0].font.size = PptPt(28)
        content_frame.paragraphs[0].font.color.rgb = PptRGBColor(255, 255, 255)

    prs.save('SOUTENANCE.pptx')
    print("✓ SOUTENANCE.pptx créé!")
else:
    print("! PowerPoint non créé (python-pptx manquant)")

print("\nFichiers créés:")
print("  - SOUTENANCE.docx")
print("  - SOUTENANCE.pptx (si python-pptx installé)")
